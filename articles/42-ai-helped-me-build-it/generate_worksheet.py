"""Build the Process Step Classification Worksheet for Article 42 --
"AI Helped Me Build It. AI Does Not Have to Run It."

Run from anywhere:
    python generate_worksheet.py

Writes, into ./templates/:
    process_step_classification_worksheet.xlsx  (PythonMuse LLC branded)
    process_step_classification_worksheet.csv   (the pre-filled example rows)

Styling follows docs/PythonMuse_BrandSample.xlsx: Deep Navy header row with
white bold text, thin Ocean Teal borders, #EDF5F5 alternating rows, a 16pt
Deep Navy title and 12pt Bright Teal subtitle, rotating brand tab colors.
All example content is synthetic.
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DEEP_NAVY = "002639"
MIDNIGHT_TEAL = "003144"
BRIGHT_TEAL = "3ABFB9"
GOLDEN_YELLOW = "FFD75E"
OCEAN_TEAL = "005F6F"
SOFT_SAGE = "91BE8E"
ALT_ROW = "EDF5F5"
WHITE = "FFFFFF"

OUT = Path(__file__).parent / "templates"

COLUMNS = [
    ("Step #", 8),
    ("Step", 30),
    ("Input", 28),
    ("Expected output", 28),
    ("Written rule?", 30),
    ("Possible failure", 30),
    ("Reviewer", 18),
    ("Evidence kept", 28),
    ("Classification", 16),
    ("Proposed approach", 28),
    ("Data restrictions", 28),
    ("Accountable person", 18),
]

# Year-over-year variance analysis -- the worked example from the article.
EXAMPLE = [
    [1, "Confirm expected files and fields are present",
     "Current- and prior-year trial balance exports",
     "Pass/fail list of expected files and columns",
     "Yes: list of expected file names and column headers",
     "File missing, column renamed, wrong entity exported",
     "Senior accountant", "Schema-check log saved with the run",
     "Rule", "Deterministic automation (Python or Power Query)",
     "Internal financial data; script runs locally, no AI at run time",
     "Controller"],
    [2, "Validate reporting periods and control totals",
     "Both trial balance files; general ledger totals",
     "Period confirmed; trial balance totals tie to the general ledger",
     "Yes: period must equal the reporting period; totals must tie",
     "Wrong period, partial export, unposted entries",
     "Senior accountant", "Control-total report saved with the run",
     "Rule", "Deterministic automation",
     "Internal financial data; no AI at run time", "Controller"],
    [3, "Calculate dollar and percentage changes",
     "Validated trial balance files",
     "Variance by account, in dollars and percent",
     "Yes: change = current year minus prior year; % = change / |prior year|",
     "Divide-by-zero when prior year is 0; sign errors on credit balances",
     "Senior accountant", "Variance file plus script version / commit",
     "Rule", "Deterministic automation",
     "Internal financial data; no AI at run time", "Controller"],
    [4, "Apply materiality thresholds",
     "Variance by account; approved thresholds",
     "List of accounts over threshold, ranked",
     "Yes: approved policy (example: over $50,000 AND over 10%)",
     "Last year's threshold still hard-coded",
     "Controller", "Threshold value and approval reference in the run log",
     "Rule", "Deterministic automation (threshold in one labeled parameter)",
     "Internal financial data; no AI at run time", "Controller"],
    [5, "Retrieve supporting transaction detail",
     "Flagged accounts; general ledger detail",
     "Transaction detail for each flagged account",
     "Yes: defined query by account and period",
     "Wrong date range; a journal source left out",
     "Staff accountant", "Query text and extracted detail saved",
     "Rule", "Query, script, or approved tool",
     "Internal financial data; no AI at run time", "Controller"],
    [6, "Draft preliminary commentary",
     "Flagged variances; summarized transaction detail",
     "Draft questions and commentary for account owners",
     "Partly: format is defined, content is interpretation",
     "Invented explanation; overconfident wording",
     "Senior accountant", "Draft labeled 'AI-assisted draft' plus reviewer edits",
     "Language", "LLM, if permitted by data policy",
     "Summarized, approved data only; no payroll or customer detail",
     "Controller"],
    [7, "Determine the actual business explanation",
     "Draft questions; responses from department leaders; support",
     "Supported explanation for each variance",
     "No: facts-and-circumstances conclusion",
     "Explanation accepted without support",
     "Controller", "Support documents linked to each explanation",
     "Judgment", "Accountant",
     "As per source documents", "Controller"],
    [8, "Approve the final analysis",
     "Completed variance analysis",
     "Signed-off analysis",
     "No: accountability",
     "Approval without reviewing the exceptions",
     "CFO", "Sign-off record with date",
     "Judgment", "Human reviewer",
     "Not applicable", "CFO"],
]

INSTRUCTIONS = [
    "How to use this worksheet",
    "",
    "1. Pick one recurring process: narrow, repetitive, and familiar. Not month-end close.",
    "2. On the 'Your Process' tab, write down each step in order, one row per step.",
    "3. For each step, fill in the input, expected output, written rule, possible failure, reviewer, and evidence kept.",
    "4. Classify the work inside the step: Rule, Pattern, Language, or Judgment (drop-down).",
    "5. Use the six questions below to choose a proposed approach.",
    "6. Ask AI to help you examine the process, not automate it. Where is undocumented judgment hiding?",
    "",
    "The 'Variance Example' tab shows a completed version for year-over-year variance analysis.",
    "",
    "Six questions for each step",
    "1. Is the step necessary and understood?  No -> simplify, document, or redesign first.",
    "2. Can we write the rule, and should the same input always give the same result?  Yes -> deterministic automation.",
    "3. Does it require interpreting language or context?  Yes -> consider an LLM, subject to data restrictions, validation, and review.",
    "4. Do reliable historical examples hold a pattern we can't write as rules?  Yes -> consider machine learning. No -> human-led or redesign.",
    "5. (Every option) Does it involve judgment, approval, restricted data, or contractual limits?  Yes -> a qualified person stays responsible.",
    "6. (Every option) What evidence will show it worked as intended?  Define validation, logging, exceptions, and sign-off before you build.",
    "",
    "Companion to the article 'AI Helped Me Build It. AI Does Not Have to Run It.' Synthetic example data only.",
    "Educational purposes only; not professional advice.",
]

thin = Side(style="thin", color=OCEAN_TEAL)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEADER_FILL = PatternFill("solid", fgColor=DEEP_NAVY)
ALT_FILL = PatternFill("solid", fgColor=ALT_ROW)
WRAP = Alignment(wrap_text=True, vertical="top")


def title_block(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=DEEP_NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Calibri", size=12, bold=True, color=BRIGHT_TEAL)
    ws.sheet_view.showGridLines = False


def table(ws, rows, blank_rows=0):
    header_row = 4
    for c, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=c, value=name)
        cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[header_row].height = 30

    all_rows = rows + [[i + len(rows) + 1] + [None] * (len(COLUMNS) - 1)
                       for i in range(blank_rows)]
    for r, values in enumerate(all_rows, start=header_row + 1):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name="Calibri", size=11, color=DEEP_NAVY,
                             bold=(c == 9 and value is not None))
            cell.border = BORDER
            cell.alignment = WRAP if c != 1 else Alignment(horizontal="center", vertical="top")
            if (r - header_row) % 2 == 0:
                cell.fill = ALT_FILL
        ws.row_dimensions[r].height = 62 if rows else 36

    last = header_row + len(all_rows)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(COLUMNS))}{last}"

    # Drop-downs keep the classification vocabulary consistent.
    dv_class = DataValidation(type="list", formula1='"Rule,Pattern,Language,Judgment"',
                              allow_blank=True)
    dv_class.add(f"I{header_row + 1}:I{last}")
    ws.add_data_validation(dv_class)

    foot = last + 2
    ws.cell(row=foot, column=1, value="PythonMuse LLC  |  www.pythonmuse.com").font = Font(
        name="Calibri", size=10, italic=True, color=OCEAN_TEAL)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build():
    wb = Workbook()

    ws = wb.active
    ws.title = "Instructions"
    ws.sheet_properties.tabColor = BRIGHT_TEAL
    title_block(ws, "Process Step Classification Worksheet",
                "Rules, patterns, language, or judgment: classify the work inside each step")
    ws.column_dimensions["A"].width = 130
    for i, line in enumerate(INSTRUCTIONS, start=4):
        cell = ws.cell(row=i, column=1, value=line)
        heading = line in ("How to use this worksheet", "Six questions for each step")
        cell.font = Font(name="Calibri", size=12 if heading else 11, bold=heading,
                         color=OCEAN_TEAL if heading else DEEP_NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=4 + len(INSTRUCTIONS) + 1, column=1,
            value="PythonMuse LLC  |  www.pythonmuse.com").font = Font(
        name="Calibri", size=10, italic=True, color=OCEAN_TEAL)

    ex = wb.create_sheet("Variance Example")
    ex.sheet_properties.tabColor = GOLDEN_YELLOW
    title_block(ex, "Example: Year-Over-Year Variance Analysis",
                "Completed model rows (synthetic). Python calculates; AI suggests questions; the accountant explains.")
    table(ex, EXAMPLE)

    yours = wb.create_sheet("Your Process")
    yours.sheet_properties.tabColor = OCEAN_TEAL
    title_block(yours, "Your Process: [name the process]",
                "One row per step. Classify each step, then choose the approach.")
    table(yours, [], blank_rows=15)

    OUT.mkdir(exist_ok=True)
    wb.save(OUT / "process_step_classification_worksheet.xlsx")

    with open(OUT / "process_step_classification_worksheet.csv", "w", newline="",
              encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow([name for name, _ in COLUMNS])
        writer.writerows(EXAMPLE)


if __name__ == "__main__":
    build()
    print("Wrote templates/process_step_classification_worksheet.xlsx and .csv")
