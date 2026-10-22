"""Tests for the Article 43 scoring demo and cost workbook.

Run from the article folder:
    python -m pytest test_score_predictions.py -q

Every expected figure below was worked out by hand from data/collection_notes.csv
and data/sample_predictions.csv BEFORE these assertions were written, using the
illustrative bands auto-route >= 0.90, confirm >= 0.60, escalate below 0.60.

  auto-route (9):  AR-001 .96, 002 .94, 003 .95, 004 .97, 005 .92, 006 .91,
                   015 .90 (boundary, inclusive), 018 .93       -> correct
                   AR-011 .93 said "other", known pricing_dispute -> WRONG
                   => 8 of 9 correct
  confirm (7):     AR-008 .88, 009 .71, 010 .83, 012 .77, 014 .69,
                   016 .89 (just under the auto line)            -> correct
                   AR-013 .64 said "other", known purchase_order -> WRONG
                   => 6 of 7 correct
  escalate (3):    AR-017 .58                                    -> correct
                   AR-007 .47 said purchase_order, known pricing -> WRONG
                   AR-019 .52 said pricing, known quantity       -> WRONG
                   => 1 of 3 correct
  no probability:  AR-020 (language-model fallback, blank)       -> correct
                   => 1 of 1
  overall:         16 of 20 = 80%

  Confusion (known -> choice):
    cash_flow         4 cash_flow
    pricing_dispute   3 pricing, 1 purchase_order (AR-007), 1 other (AR-011)
    purchase_order    3 purchase_order, 1 other (AR-013)
    quantity_dispute  3 quantity, 1 pricing (AR-019)
    other             3 other

  The only high-confidence error is AR-011, the note with an instruction
  planted in it. It is deliberate: do not "fix" it out of the sample data.

  Cost per acceptable decision (10,000 notes, $60/hour, $50 per uncaught error):
    Decision model + bands: 0.50 + 8,000 + 3,500 = 11,500.50 / 9,930 = 1.1582
    Cheapest model, half to a person: 0.25 + 30,000 + 5,000 = 35,000.25 / 9,900
                                      = 3.5354
    LLM paragraph, every item read: 40 + 15,000 + 0 = 15,040 / 10,000 = 1.504
    All manual: 30,000 / 10,000 = 3.00
"""

import io
import subprocess
import sys
from pathlib import Path

import pytest

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))

import score_predictions as sp  # noqa: E402
import generate_worksheet as gw  # noqa: E402


@pytest.fixture(scope="module")
def results():
    return sp.score(sp.read_rows(sp.DEFAULT_LABELS),
                    sp.read_rows(sp.DEFAULT_PREDICTIONS))


def test_overall_accuracy(results):
    assert results["total"] == 20
    assert results["correct"] == 16
    assert results["accuracy"] == pytest.approx(0.80)


@pytest.mark.parametrize("band, items, correct", [
    (sp.AUTO, 9, 8),
    (sp.CONFIRM, 7, 6),
    (sp.ESCALATE, 3, 1),
    (sp.NO_PROB, 1, 1),
])
def test_accuracy_by_band(results, band, items, correct):
    b = results["by_band"][band]
    assert (b["items"], b["correct"]) == (items, correct)


def test_boundaries_are_inclusive_at_the_lower_edge(results):
    band = {i["item_id"]: i["band"] for i in results["items"]}
    assert band["AR-015"] == sp.AUTO        # exactly 0.90
    assert band["AR-016"] == sp.CONFIRM     # 0.89, just under
    assert band["AR-017"] == sp.ESCALATE    # 0.58
    assert band["AR-020"] == sp.NO_PROB     # blank: no probability


def test_the_47_42_11_example_is_escalated_not_auto_routed(results):
    item = next(i for i in results["items"] if i["item_id"] == "AR-007")
    assert item["probability"] == pytest.approx(0.47)
    assert item["band"] == sp.ESCALATE
    assert not item["correct"]


def test_only_high_confidence_error_is_the_planted_instruction(results):
    errors = results["high_confidence_errors"]
    assert [e["item_id"] for e in errors] == ["AR-011"]
    assert errors[0]["choice"] == "other"
    assert errors[0]["known"] == "pricing_dispute"


def test_confusion_table(results):
    c = results["confusion"]
    assert c["cash_flow"]["cash_flow"] == 4
    assert c["pricing_dispute"] == {"cash_flow": 0, "pricing_dispute": 3,
                                    "purchase_order": 1, "quantity_dispute": 0,
                                    "other": 1}
    assert c["purchase_order"]["purchase_order"] == 3
    assert c["purchase_order"]["other"] == 1
    assert c["quantity_dispute"]["quantity_dispute"] == 3
    assert c["quantity_dispute"]["pricing_dispute"] == 1
    assert c["other"]["other"] == 3
    assert sum(sum(row.values()) for row in c.values()) == 20
    assert sum(c[k][k] for k in sp.ALLOWED_CHOICES) == 16


def test_tighter_bands_move_the_planted_error_out_of_auto_route():
    r = sp.score(sp.read_rows(sp.DEFAULT_LABELS),
                 sp.read_rows(sp.DEFAULT_PREDICTIONS),
                 auto_at=0.95, confirm_at=0.60)
    # Only AR-001 (.96), AR-003 (.95), AR-004 (.97) clear 0.95.
    assert r["by_band"][sp.AUTO]["items"] == 3
    assert r["high_confidence_errors"] == []


def test_choice_outside_the_allowed_list_is_rejected():
    labels = [{"item_id": "X1", "known_outcome": "other"}]
    preds = [{"item_id": "X1", "model_choice": "probably fine", "probability": "0.99"}]
    with pytest.raises(sp.ScoringError, match="not an allowed choice"):
        sp.score(labels, preds)


def test_missing_and_unknown_items_are_rejected():
    labels = [{"item_id": "X1", "known_outcome": "other"},
              {"item_id": "X2", "known_outcome": "other"}]
    with pytest.raises(sp.ScoringError, match="no prediction for: X2"):
        sp.score(labels, [{"item_id": "X1", "model_choice": "other", "probability": ""}])
    with pytest.raises(sp.ScoringError, match="no known outcome"):
        sp.score(labels, [{"item_id": "X9", "model_choice": "other", "probability": ""}])


def test_bad_probability_is_rejected():
    labels = [{"item_id": "X1", "known_outcome": "other"}]
    with pytest.raises(sp.ScoringError, match="outside 0 to 1"):
        sp.score(labels, [{"item_id": "X1", "model_choice": "other", "probability": "96"}])


def test_report_and_command_line_run():
    out = sp.report(sp.score(sp.read_rows(sp.DEFAULT_LABELS),
                             sp.read_rows(sp.DEFAULT_PREDICTIONS)), out=io.StringIO())
    text = out.getvalue()
    assert "16 correct" in text and "AR-011" in text
    run = subprocess.run([sys.executable, str(HERE / "score_predictions.py")],
                         capture_output=True, text=True, check=True)
    assert "AR-011: model said other at 93%" in run.stdout


# ---------------------------------------------------------------------------
# Cost per acceptable decision (the Python mirror of the workbook formulas)
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("name, total, acceptable, per_decision", [
    ("Decision model + bands", 11500.50, 9930, 1.1582),
    ("Cheapest model, loose bands", 35000.25, 9900, 3.5354),
    ("LLM paragraph, every item read", 15040.00, 10000, 1.5040),
    ("All manual", 30000.00, 10000, 3.0000),
])
def test_cost_per_acceptable_decision(name, total, acceptable, per_decision):
    scenario = next(s for s in gw.SCENARIOS if s["name"] == name)
    r = gw.cost_per_acceptable_decision(scenario)
    assert r["total_cost"] == pytest.approx(total)
    assert r["acceptable"] == pytest.approx(acceptable)
    assert r["per_decision"] == pytest.approx(per_decision, abs=1e-4)


def test_workbooks_have_live_formulas(tmp_path):
    from openpyxl import load_workbook
    gw.build(tmp_path)
    cost = load_workbook(tmp_path / "cost_per_acceptable_decision.xlsx")["Compare"]
    formulas = [c.value for row in cost.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) >= 24      # 6 calculated rows x 4 scenarios
    pilot = load_workbook(tmp_path / "decision_model_pilot_worksheet.xlsx")["AR Pilot Example"]
    correct_col = [pilot.cell(row=r, column=7).value for r in range(5, 25)]
    assert all(str(v).startswith("=IF(") for v in correct_col)
